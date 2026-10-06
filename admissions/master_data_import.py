from __future__ import annotations

from collections import Counter
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from django.db import transaction

from .models import (
    Institution,
    Program,
    ProgramFeePlan,
    ProgramFeeInstallment,
    ProgramInternalFinance,
)

REQUIRED_SHEETS = {
    "Institutions",
    "Offerings",
    "Fee_Plans",
    "Installments",
    "Additional_Fees",
    "Course_Streams",
    "Rules_Notes",
    "Data_Quality",
}

REVIEW_ONLY_SHEETS = {"Additional_Fees", "Course_Streams", "Rules_Notes", "Data_Quality"}


def clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def decimal_value(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    try:
        return Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        return None


def integer_value(value: Any) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(float(value))
    except (ValueError, TypeError):
        return None


def source_notes(row: dict[str, Any]) -> str:
    parts = []
    for key in ("source_file", "source_location", "notes"):
        value = clean(row.get(key))
        if value:
            parts.append(f"{key}: {value}")
    return " | ".join(parts)


def read_workbook(workbook_path: str | Path) -> dict[str, list[dict[str, Any]]]:
    try:
        import openpyxl
    except ImportError as exc:
        raise ValueError("openpyxl is required for Excel imports.") from exc

    wb = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    missing = sorted(REQUIRED_SHEETS - set(wb.sheetnames))
    if missing:
        raise ValueError("Missing required sheets: " + ", ".join(missing))

    data: dict[str, list[dict[str, Any]]] = {}
    for sheet in REQUIRED_SHEETS:
        ws = wb[sheet]
        rows = ws.iter_rows(values_only=True)
        headers = [clean(v) for v in next(rows, ())]
        if not headers or any(not h for h in headers):
            raise ValueError(f"Sheet {sheet} has an invalid header row.")
        records = []
        for row_number, row in enumerate(rows, start=2):
            if not any(clean(v) for v in row):
                continue
            record = {headers[i]: row[i] if i < len(row) else None for i in range(len(headers)) if headers[i]}
            record["_row_number"] = row_number
            records.append(record)
        data[sheet] = records
    return data


def _duplicates(values: list[str]) -> list[str]:
    counts = Counter(v for v in values if v)
    return sorted(v for v, count in counts.items() if count > 1)


def level_from_name(name: str) -> str:
    n = clean(name).upper()
    if "CERTIFICATE" in n or n.startswith("CERT"):
        return "CERTIFICATE"
    if "DIPLOMA" in n or n.startswith("DCA") or "PGD" in n:
        return "DIPLOMA"
    if n.startswith(("B.", "B ", "BBA", "BCA", "BCOM", "B.COM", "BA", "B.SC", "BSC", "LLB")) or "BACHELOR" in n:
        return "UG"
    if n.startswith(("M.", "M ", "MBA", "MCA", "M.COM", "MA", "M.SC", "MSC", "LLM")) or "MASTER" in n:
        return "PG"
    return "OTHER"


def study_mode(value: Any) -> str:
    value = clean(value).upper()
    return value if value in {"DISTANCE", "ONLINE", "REGULAR", "HYBRID", "OTHER"} else "OTHER"


def duration(row: dict[str, Any]) -> tuple[Decimal | None, int | None]:
    value = decimal_value(row.get("duration_value"))
    unit = clean(row.get("duration_unit")).lower()
    semesters = integer_value(row.get("total_semesters"))
    if unit == "semester" and value is not None:
        return None, int(value)
    if unit == "year" and value is not None:
        return value, semesters
    if unit == "month" and value is not None:
        return value / Decimal("12"), semesters
    return None, semesters


def student_plan_name(row: dict[str, Any]) -> str:
    label = clean(row.get("fee_plan_label")) or "Standard"
    frequency = clean(row.get("course_fee_frequency")) or "STANDARD"
    return f"{label} — {frequency}"[:120]



def _has_value(value: Any) -> bool:
    """Return True when the source value contains meaningful data."""
    return clean(value) != ""


def program_data_status(row: dict[str, Any]) -> str:
    """
    Determine the initial verification state for an imported Program.

    VERIFIED is intentionally never assigned by the importer.
    Human verification is required before a record becomes VERIFIED.
    """
    status = clean(row.get("status")).upper()

    # Source explicitly identifies a non-ready record.
    if status and status != "READY":
        return Program.DataStatus.NEEDS_REVIEW

    important_fields = (
        row.get("course_name"),
        row.get("study_mode"),
        row.get("eligibility"),
        row.get("duration_value"),
        row.get("duration_unit"),
        row.get("total_semesters"),
    )

    missing_count = sum(not _has_value(value) for value in important_fields)

    if missing_count:
        return Program.DataStatus.PARTIAL

    return Program.DataStatus.NEEDS_REVIEW


def fee_plan_data_status(row: dict[str, Any]) -> str:
    """
    Determine the initial verification state for an imported fee plan.

    VERIFIED is intentionally never assigned by the importer.
    """
    source_status = clean(row.get("data_status")).upper()

    if source_status and source_status != "READY":
        return ProgramFeePlan.DataStatus.NEEDS_REVIEW

    important_fields = (
        row.get("total_fee"),
        row.get("admission_fee"),
        row.get("exam_fee"),
        row.get("course_fee"),
        row.get("course_fee_frequency"),
    )

    missing_count = sum(not _has_value(value) for value in important_fields)

    if missing_count:
        return ProgramFeePlan.DataStatus.PARTIAL

    return ProgramFeePlan.DataStatus.NEEDS_REVIEW


def nullable_decimal(value: Any) -> Decimal | None:
    """
    Parse a decimal while preserving missing source values as NULL.

    Invalid non-empty values also become None; validation/review status
    is responsible for ensuring such records are reviewed rather than
    silently represented as zero.
    """
    if value is None or clean(value) == "":
        return None

    return decimal_value(value)


def source_verification_note(row: dict[str, Any]) -> str:
    """
    Preserve useful source metadata in the verification notes without
    marking the record as verified.
    """
    notes = source_notes(row)
    return (
        "Imported from academic master workbook. "
        "Record requires human verification before it can be marked VERIFIED."
        + (f" {notes}" if notes else "")
    )


def validate(data: dict[str, list[dict[str, Any]]]) -> dict[str, Any]:
    errors: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []

    institutions = data["Institutions"]
    offerings = data["Offerings"]
    fee_plans = data["Fee_Plans"]
    installments = data["Installments"]

    institution_codes = {clean(r.get("institution_code")).upper() for r in institutions if clean(r.get("institution_code"))}
    offering_keys = {clean(r.get("offering_key")) for r in offerings if clean(r.get("offering_key"))}
    fee_plan_keys = {clean(r.get("fee_plan_key")) for r in fee_plans if clean(r.get("fee_plan_key"))}

    for label, values in (
        ("institution_code", [clean(r.get("institution_code")).upper() for r in institutions]),
        ("offering_key", [clean(r.get("offering_key")) for r in offerings]),
        ("fee_plan_key", [clean(r.get("fee_plan_key")) for r in fee_plans]),
    ):
        for duplicate in _duplicates(values):
            errors.append({"code": "DUPLICATE_KEY", "message": f"Duplicate {label}: {duplicate}"})

    for row in offerings:
        code = clean(row.get("institution_code")).upper()
        key = clean(row.get("offering_key"))
        if not key:
            errors.append({"code": "MISSING_OFFERING_KEY", "row": row.get("_row_number"), "message": "Offering key is required."})
        if code not in institution_codes:
            errors.append({"code": "UNKNOWN_INSTITUTION", "row": row.get("_row_number"), "message": f"Offering {key} references unknown institution {code}."})

    for row in fee_plans:
        key = clean(row.get("fee_plan_key"))
        offering = clean(row.get("offering_key"))
        category = clean(row.get("fee_category")).upper()
        if offering not in offering_keys:
            errors.append({"code": "UNKNOWN_OFFERING", "row": row.get("_row_number"), "message": f"Fee plan {key} references unknown offering {offering}."})
        if category not in {"STUDENT_FEE", "CENTER_FEE"}:
            errors.append({"code": "UNKNOWN_FEE_CATEGORY", "row": row.get("_row_number"), "message": f"Fee plan {key} has unsupported category {category}."})

    for row in installments:
        key = clean(row.get("fee_plan_key"))
        if key not in fee_plan_keys:
            errors.append({"code": "UNKNOWN_FEE_PLAN", "row": row.get("_row_number"), "message": f"Installment references unknown fee_plan_key {key}."})
        if integer_value(row.get("installment_no")) is None:
            errors.append({"code": "INVALID_INSTALLMENT", "row": row.get("_row_number"), "message": f"Installment number is invalid for {key}."})

    for row in data["Data_Quality"]:
        description = clean(row.get("description"))
        if description:
            severity = clean(row.get("severity")).upper() or "INFO"
            warnings.append({
                "code": clean(row.get("issue_type")) or "DATA_QUALITY",
                "severity": severity,
                "institution_code": clean(row.get("institution_code")),
                "message": description,
                "action_required": clean(row.get("action_required")),
            })

    for row in offerings:
        status = clean(row.get("status")).upper()
        if status and status != "READY":
            warnings.append({
                "code": status,
                "severity": "MEDIUM",
                "institution_code": clean(row.get("institution_code")),
                "message": f"Offering {clean(row.get('offering_key'))} has source status {status}.",
            })

    for row in fee_plans:
        status = clean(row.get("data_status")).upper()
        if status and status != "READY":
            warnings.append({
                "code": status,
                "severity": "MEDIUM",
                "institution_code": clean(row.get("institution_code")),
                "message": f"Fee plan {clean(row.get('fee_plan_key'))} has source status {status}.",
            })

    existing_institution_codes = set(
        Institution.objects.filter(code__in=list(institution_codes)).values_list("code", flat=True)
    )

    existing_programs = list(
        Program.objects.filter(code__in=list(offering_keys))
        .select_related("institution")
    )
    existing_program_code_counts = Counter(p.code for p in existing_programs)

    for code, count in sorted(existing_program_code_counts.items()):
        if count > 1:
            errors.append(
                {
                    "code": "AMBIGUOUS_EXISTING_PROGRAM_CODE",
                    "message": (
                        f"Existing database contains {count} programs with code {code!r}. "
                        "Import is blocked until the duplicate program codes are resolved."
                    ),
                    "offering_key": code,
                }
            )

    program_lookup = {
        p.code: p
        for p in existing_programs
        if existing_program_code_counts[p.code] == 1
    }

    workbook_institution_by_code = {
        clean(row.get("institution_code")): clean(row.get("university_name"))
        for row in institutions
        if clean(row.get("institution_code"))
    }

    for code, program in sorted(program_lookup.items()):
        workbook_institution_code = ""
        for row in offerings:
            if clean(row.get("offering_key")) == code:
                workbook_institution_code = clean(row.get("institution_code"))
                break

        if (
            workbook_institution_code
            and program.institution.code
            and program.institution.code != workbook_institution_code
        ):
            errors.append(
                {
                    "code": "PROGRAM_INSTITUTION_CONFLICT",
                    "message": (
                        f"Existing program {code!r} belongs to institution "
                        f"{program.institution.code!r}, but the workbook assigns it to "
                        f"{workbook_institution_code!r}. Import is blocked for this program."
                    ),
                    "offering_key": code,
                }
            )

    existing_fee_plan_keys: set[str] = set()
    existing_fee_plans = {
        obj.import_key: obj
        for obj in ProgramFeePlan.objects.filter(
            import_key__in=[
                clean(row.get("fee_plan_key"))
                for row in fee_plans
                if clean(row.get("fee_category")).upper() == "STUDENT_FEE"
                and clean(row.get("fee_plan_key"))
            ]
        ).select_related("program")
        if obj.import_key
    }

    for row in fee_plans:
        if clean(row.get("fee_category")).upper() != "STUDENT_FEE":
            continue

        key = clean(row.get("fee_plan_key"))
        if not key:
            continue

        existing = existing_fee_plans.get(key)
        if existing:
            workbook_program_code = clean(row.get("offering_key"))
            if existing.program.code != workbook_program_code:
                errors.append(
                    {
                        "code": "FEE_PLAN_IMPORT_KEY_CONFLICT",
                        "message": (
                            f"Fee plan import_key {key!r} already belongs to program "
                            f"{existing.program.code!r}, but the workbook assigns it to "
                            f"{workbook_program_code!r}. Import is blocked."
                        ),
                        "fee_plan_key": key,
                    }
                )
            else:
                existing_fee_plan_keys.add(key)

    new_institutions = len(institution_codes - existing_institution_codes)
    new_programs = len(offering_keys - set(program_lookup.keys()))
    update_institutions = len(existing_institution_codes)
    update_programs = len(program_lookup)
    new_student_plans = sum(
        1
        for row in fee_plans
        if clean(row.get("fee_category")).upper() == "STUDENT_FEE"
        and clean(row.get("fee_plan_key"))
        and clean(row.get("fee_plan_key")) not in existing_fee_plan_keys
    )
    update_student_plans = sum(
        1
        for row in fee_plans
        if clean(row.get("fee_category")).upper() == "STUDENT_FEE"
        and clean(row.get("fee_plan_key")) in existing_fee_plan_keys
    )

    counts = {
        "institutions": len(institutions),
        "offerings": len(offerings),
        "fee_plans": len(fee_plans),
        "student_fee_plans": sum(1 for r in fee_plans if clean(r.get("fee_category")).upper() == "STUDENT_FEE"),
        "center_fee_plans": sum(1 for r in fee_plans if clean(r.get("fee_category")).upper() == "CENTER_FEE"),
        "installments": len(installments),
        "additional_fees": len(data["Additional_Fees"]),
        "course_streams": len(data["Course_Streams"]),
        "rules_notes": len(data["Rules_Notes"]),
        "data_quality": len(data["Data_Quality"]),
    }

    return {
        "valid": not errors,
        "errors": errors,
        "warnings": warnings,
        "requires_warning_acknowledgement": any(w.get("severity") == "HIGH" for w in warnings),
        "counts": counts,
        "changes": {
            "new_institutions": new_institutions,
            "update_institutions": update_institutions,
            "new_programs": new_programs,
            "update_programs": update_programs,
            "new_student_fee_plans": new_student_plans,
            "update_student_fee_plans": update_student_plans,
            "center_finance_records": counts["center_fee_plans"],
            "installments": counts["installments"],
        },
        "review_only": {
            "additional_fees": counts["additional_fees"],
            "course_streams": counts["course_streams"],
            "rules_notes": counts["rules_notes"],
            "data_quality": counts["data_quality"],
        },
    }


def import_data(data: dict[str, list[dict[str, Any]]]) -> dict[str, int]:
    counts: Counter[str] = Counter()
    institution_map: dict[str, Institution] = {}
    program_map: dict[str, Program] = {}
    fee_plan_map: dict[str, ProgramFeePlan] = {}

    with transaction.atomic():
        for row in data["Institutions"]:
            code = clean(row.get("institution_code")).upper()
            if not code:
                continue
            obj, created = Institution.objects.update_or_create(
                code=code,
                defaults={
                    "name": clean(row.get("university_name")) or code,
                    "short_name": code,
                    "is_active": clean(row.get("status")).upper() != "INACTIVE",
                    "notes": source_notes(row),
                },
            )
            institution_map[code] = obj
            counts[f"institutions_{'created' if created else 'updated'}"] += 1

        for row in data["Offerings"]:
            key = clean(row.get("offering_key"))
            institution = institution_map.get(clean(row.get("institution_code")).upper())
            if not key or not institution:
                continue
            years, semesters = duration(row)
            obj, created = Program.objects.update_or_create(
                institution=institution,
                code=key,
                defaults={
                    "name": clean(row.get("course_name")) or key,
                    "level": level_from_name(row.get("course_name")),
                    "duration_years": years,
                    "duration_semesters": semesters,
                    "study_mode": study_mode(row.get("study_mode")),
                    "specialization": clean(row.get("specializations_raw")),
                    "eligibility_text": clean(row.get("eligibility")),
                    "minimum_qualification": "UNSPECIFIED",
                    "required_stream": "",
                    "eligibility_review_required": clean(row.get("status")).upper() not in {"", "READY"},
                    "is_credit_transfer_available": False,
                    "is_active": clean(row.get("status")).upper() != "INACTIVE",
                    "data_status": program_data_status(row),
                    "verified_at": None,
                    "verified_by": None,
                    "verification_notes": source_verification_note(row),
                    "notes": source_notes(row),
                },
            )
            program_map[key] = obj
            counts[f"programs_{'created' if created else 'updated'}"] += 1

        for row in data["Fee_Plans"]:
            key = clean(row.get("fee_plan_key"))
            offering = clean(row.get("offering_key"))
            category = clean(row.get("fee_category")).upper()
            program = program_map.get(offering)
            if not key or not program:
                continue

            if category == "CENTER_FEE":
                center = decimal_value(row.get("centre_fee") or row.get("total_fee") or row.get("course_fee"))
                if center is not None:
                    ProgramInternalFinance.objects.update_or_create(
                        program=program,
                        defaults={"center_fee": center, "notes": source_notes(row)},
                    )
                    counts["center_finance_processed"] += 1
                continue

            if category != "STUDENT_FEE":
                continue

            obj, created = ProgramFeePlan.objects.update_or_create(
                import_key=key,
                defaults={
                    "program": program,
                    "name": student_plan_name(row),
                    "student_total_fee": nullable_decimal(row.get("total_fee")),
                    "registration_fee": nullable_decimal(row.get("admission_fee")),
                    "exam_fee": nullable_decimal(row.get("exam_fee")),
                    "other_fee": None,
                    "is_active": clean(row.get("data_status")).upper() != "INACTIVE",
                    "data_status": fee_plan_data_status(row),
                    "verified_at": None,
                    "verified_by": None,
                    "verification_notes": source_verification_note(row),
                    "notes": source_notes(row),
                },
            )
            fee_plan_map[key] = obj
            counts[f"fee_plans_{'created' if created else 'updated'}"] += 1

        for row in data["Installments"]:
            plan = fee_plan_map.get(clean(row.get("fee_plan_key")))
            number = integer_value(row.get("installment_no"))
            if not plan or number is None:
                continue
            ProgramFeeInstallment.objects.update_or_create(
                fee_plan=plan,
                installment_number=number,
                defaults={
                    "label": clean(row.get("installment_name")) or f"Installment {number}",
                    "amount": nullable_decimal(row.get("amount")),
                    "due_stage": clean(row.get("payment_trigger")),
                    "notes": clean(row.get("notes")),
                },
            )
            counts["installments_processed"] += 1

    counts["additional_fees_review_only"] = len(data["Additional_Fees"])
    counts["course_streams_review_only"] = len(data["Course_Streams"])
    counts["rules_notes_review_only"] = len(data["Rules_Notes"])
    counts["data_quality_rows"] = len(data["Data_Quality"])
    return dict(counts)
