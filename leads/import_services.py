import csv
import io
import re

from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from .models import (
    Lead,
    LeadActivity,
    LeadImportBatch,
    LeadImportRow,
)


HEADER_ALIASES = {
    "name": {
        "name",
        "full name",
        "fullname",
        "customer name",
        "lead name",
    },
    "phone_number": {
        "phone",
        "phone number",
        "phone_number",
        "mobile",
        "mobile number",
        "mobilenumber",
        "contact",
        "contact number",
        "whatsapp",
        "whatsapp number",
    },
    "email": {
        "email",
        "email address",
        "mail",
    },
    "city": {
        "city",
        "location",
    },
    "state": {
        "state",
    },
    "interested_course": {
        "course",
        "interested course",
        "course interested",
        "program",
        "programme",
    },
}


def normalize_header(value):
    if value is None:
        return ""

    value = str(value).strip().lower()
    value = re.sub(r"[_\-]+", " ", value)
    value = re.sub(r"\s+", " ", value)

    return value


def identify_field(header):
    normalized = normalize_header(header)

    for field, aliases in HEADER_ALIASES.items():
        if normalized in aliases:
            return field

    return None


def normalize_phone(value):
    if value is None:
        return ""

    if isinstance(value, float):
        if value.is_integer():
            value = str(int(value))

    value = str(value).strip()

    if value.endswith(".0"):
        value = value[:-2]

    digits = re.sub(r"\D", "", value)

    # Indian phone normalization.
    if len(digits) == 12 and digits.startswith("91"):
        digits = digits[2:]

    elif len(digits) == 11 and digits.startswith("0"):
        digits = digits[1:]

    return digits


def validate_phone(phone):
    """
    Current BEOIS marketing import rule:
    Indian mobile numbers are normalized to 10 digits.
    """

    if not phone:
        return False

    return (
        len(phone) == 10
        and phone.isdigit()
    )


def read_csv_file(uploaded_file):
    raw = uploaded_file.read()

    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")

    reader = csv.reader(
        io.StringIO(text)
    )

    return list(reader)


def read_excel_file(uploaded_file):
    workbook = load_workbook(
        uploaded_file,
        read_only=True,
        data_only=True,
    )

    sheet = workbook.active

    rows = [
        list(row)
        for row in sheet.iter_rows(
            values_only=True
        )
    ]

    workbook.close()

    return rows


def read_uploaded_file(uploaded_file):
    file_name = uploaded_file.name.lower()

    if file_name.endswith(".csv"):
        return read_csv_file(uploaded_file)

    if file_name.endswith(".xlsx"):
        return read_excel_file(uploaded_file)

    raise ValueError(
        "Only .xlsx and .csv files are supported."
    )


def build_column_map(headers):
    mapping = {}

    for index, header in enumerate(headers):
        field = identify_field(header)

        if field and field not in mapping:
            mapping[field] = index

    return mapping


def cell_value(row, mapping, field):
    index = mapping.get(field)

    if index is None:
        return ""

    if index >= len(row):
        return ""

    value = row[index]

    if value is None:
        return ""

    return str(value).strip()


@transaction.atomic
def create_import_preview(
    *,
    uploaded_file,
    source,
    campaign,
    default_vertical,
    default_channel,
    uploaded_by,
):
    if default_channel == Lead.Channel.PARTNER:
        raise ValueError(
            "Partner imports are not supported in "
            "Marketing Lead Import V1. "
            "Use BEST Direct for social-media leads."
        )

    rows = read_uploaded_file(
        uploaded_file
    )

    if not rows:
        raise ValueError(
            "The uploaded file is empty."
        )

    headers = rows[0]
    data_rows = rows[1:]

    mapping = build_column_map(
        headers
    )

    missing = []

    if "name" not in mapping:
        missing.append("Name")

    if "phone_number" not in mapping:
        missing.append("Phone Number")

    if missing:
        raise ValueError(
            "Required column(s) not found: "
            + ", ".join(missing)
        )

    batch = LeadImportBatch.objects.create(
        file_name=uploaded_file.name,
        source=source,
        campaign=campaign,
        default_vertical=default_vertical,
        default_channel=default_channel,
        uploaded_by=uploaded_by,
    )

    valid_count = 0
    duplicate_count = 0
    invalid_count = 0
    total_count = 0

    phones_seen_in_file = set()

    for excel_row_number, row in enumerate(
        data_rows,
        start=2,
    ):
        # Ignore completely empty spreadsheet rows.
        if not any(
            value not in (None, "")
            for value in row
        ):
            continue

        total_count += 1

        name = cell_value(
            row,
            mapping,
            "name",
        )

        raw_phone = cell_value(
            row,
            mapping,
            "phone_number",
        )

        phone = normalize_phone(
            raw_phone
        )

        email = cell_value(
            row,
            mapping,
            "email",
        )

        city = cell_value(
            row,
            mapping,
            "city",
        )

        state = cell_value(
            row,
            mapping,
            "state",
        )

        interested_course = cell_value(
            row,
            mapping,
            "interested_course",
        )

        row_status = LeadImportRow.Status.VALID
        error_message = ""
        existing_lead = None

        if not name:
            row_status = (
                LeadImportRow.Status.INVALID
            )
            error_message = (
                "Name is required."
            )

        elif not validate_phone(phone):
            row_status = (
                LeadImportRow.Status.INVALID
            )
            error_message = (
                "Phone number must contain "
                "a valid 10-digit mobile number."
            )

        elif phone in phones_seen_in_file:
            row_status = (
                LeadImportRow.Status.DUPLICATE
            )
            error_message = (
                "Duplicate phone number "
                "inside this uploaded file."
            )

        else:
            existing_lead = (
                Lead.objects
                .filter(
                    phone_number=phone
                )
                .order_by(
                    "-created_at"
                )
                .first()
            )

            if existing_lead:
                row_status = (
                    LeadImportRow
                    .Status
                    .DUPLICATE
                )
                error_message = (
                    "Phone number already exists "
                    f"as {existing_lead.lead_id}."
                )

        if phone:
            phones_seen_in_file.add(
                phone
            )

        if row_status == LeadImportRow.Status.VALID:
            valid_count += 1

        elif (
            row_status
            == LeadImportRow.Status.DUPLICATE
        ):
            duplicate_count += 1

        else:
            invalid_count += 1

        LeadImportRow.objects.create(
            batch=batch,
            row_number=excel_row_number,
            name=name,
            phone_number=phone,
            email=email,
            city=city,
            state=state,
            interested_course=interested_course,
            status=row_status,
            error_message=error_message,
            existing_lead=existing_lead,
        )

    batch.total_rows = total_count
    batch.valid_rows = valid_count
    batch.duplicate_rows = duplicate_count
    batch.invalid_rows = invalid_count
    batch.status = (
        LeadImportBatch.Status.VALIDATED
    )

    batch.save(
        update_fields=[
            "total_rows",
            "valid_rows",
            "duplicate_rows",
            "invalid_rows",
            "status",
        ]
    )

    return batch


@transaction.atomic
def confirm_import(
    *,
    batch,
    performed_by,
):
    batch = (
        LeadImportBatch.objects
        .select_for_update()
        .get(pk=batch.pk)
    )

    if (
        batch.status
        == LeadImportBatch.Status.COMPLETED
    ):
        raise ValueError(
            "This batch has already been imported."
        )

    if (
        batch.status
        != LeadImportBatch.Status.VALIDATED
    ):
        raise ValueError(
            "Only a validated batch can be imported."
        )

    batch.status = (
        LeadImportBatch.Status.IMPORTING
    )
    batch.save(
        update_fields=["status"]
    )

    imported_count = 0

    rows = (
        batch.rows
        .select_for_update()
        .filter(
            status=LeadImportRow.Status.VALID
        )
        .order_by("row_number")
    )

    for row in rows:

        # Re-check because another lead could have been created
        # after preview and before confirmation.
        existing = (
            Lead.objects
            .filter(
                phone_number=row.phone_number
            )
            .order_by("-created_at")
            .first()
        )

        if existing:
            row.status = (
                LeadImportRow.Status.DUPLICATE
            )
            row.existing_lead = existing
            row.error_message = (
                "Phone number already existed "
                "when the import was confirmed."
            )

            row.save(
                update_fields=[
                    "status",
                    "existing_lead",
                    "error_message",
                ]
            )

            continue

        lead = Lead(
            name=row.name,
            phone_number=row.phone_number,
            email=row.email,
            city=row.city,
            state=row.state,
            interested_course=(
                row.interested_course
            ),
            source=batch.source,
            campaign=batch.campaign,
            vertical=batch.default_vertical,
            channel=batch.default_channel,
            status=Lead.Status.NEW,
            created_by=performed_by,
        )

        lead.full_clean()
        lead.save()

        LeadActivity.objects.create(
            lead=lead,
            activity_type=(
                LeadActivity
                .ActivityType
                .CREATED
            ),
            description=(
                "Lead imported from marketing "
                f"batch: {batch.file_name}."
            ),
            performed_by=performed_by,
        )

        row.status = (
            LeadImportRow.Status.IMPORTED
        )
        row.imported_lead = lead
        row.error_message = ""

        row.save(
            update_fields=[
                "status",
                "imported_lead",
                "error_message",
            ]
        )

        imported_count += 1

    batch.imported_rows = imported_count

    # Recalculate because a VALID row can become DUPLICATE
    # between preview and confirmation.
    batch.duplicate_rows = (
        batch.rows.filter(
            status=LeadImportRow.Status.DUPLICATE
        ).count()
    )

    batch.valid_rows = (
        batch.rows.filter(
            status=LeadImportRow.Status.VALID
        ).count()
    )

    batch.status = (
        LeadImportBatch.Status.COMPLETED
    )

    batch.completed_at = timezone.now()

    batch.save(
        update_fields=[
            "imported_rows",
            "duplicate_rows",
            "valid_rows",
            "status",
            "completed_at",
        ]
    )

    return batch